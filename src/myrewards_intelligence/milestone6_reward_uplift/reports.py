"""Word and markdown deliverables for Milestone 6."""

from __future__ import annotations

import json
import shutil
from pathlib import Path
from typing import Any

import pandas as pd
from docx import Document
from docx.shared import Inches

from ..core.config import Settings, resolve_path
from ..core.reporting import add_bullets, add_table, new_doc
from .domain_resolutions import DOMAIN_RESOLUTIONS, OBSERVATION_WINDOW_DAYS
from .evidence import TreatmentSource


def _source_basis(results: dict[str, Any]) -> str:
    return str(results.get("treatment_source_basis") or TreatmentSource.RESEARCH_SIMULATED.value)


def _is_mock_track_r(results: dict[str, Any]) -> bool:
    return _source_basis(results) == TreatmentSource.RESEARCH_GENERATED_LEDGER.value


def _title_suffix(results: dict[str, Any]) -> str:
    if _is_mock_track_r(results):
        return "Track R — Research-generated IncentiveRewards mock"
    if _source_basis(results) == TreatmentSource.INCOMM_ISSUED.value:
        return "Track R — InComm-issued"
    return "Track S — Research-simulated"


def _m7_notice(results: dict[str, Any]) -> str:
    if results.get("m7_may_consume") and _source_basis(results) == TreatmentSource.INCOMM_ISSUED.value:
        return "TreatmentSourceBasis = InComm-issued. Milestone 7 may consume these scores."
    return (
        f"TreatmentSourceBasis = {_source_basis(results)}. "
        "Milestone 7 must not consume these scores as production lift."
    )


def _notice(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True


def _add_image(doc: Document, path: Path | None) -> None:
    if path is None:
        return
    image = Path(path)
    if image.exists():
        doc.add_picture(str(image), width=Inches(6.2))


def build_method_card(results: dict[str, Any]) -> Document:
    doc = new_doc(
        f"Reward Uplift Model v1 ({_title_suffix(results)})",
        "Milestone 6 — InComm MyRewards / KSU mHealth Lab",
    )
    _notice(doc, _m7_notice(results))
    doc.add_heading("Estimand", level=1)
    if _is_mock_track_r(results):
        doc.add_paragraph(
            "τ(x) = P(Y=1 | T=1, X=x) − P(Y=1 | T=0, X=x). T is a research-generated "
            "IncentiveRewards issuance (Bernoulli among gaps still open at outreach). "
            f"Y is InComm-sourced closure in (issued_at, issued_at+"
            f"{OBSERVATION_WINDOW_DAYS}]. Assignment is independent of later Closed. "
            "Historical closure was not caused by mock T — identified CATE is a "
            "pipeline / negative-control result, not partner-facing lift."
        )
    else:
        doc.add_paragraph(
            "τ(x) = P(Y=1 | T=1, X=x) − P(Y=1 | T=0, X=x), where Y is InComm-sourced "
            "care-gap closure and T is issued (not redeemed). Observation window "
            f"default is {OBSERVATION_WINDOW_DAYS} days; the current extract has no "
            "IssuedAtUtc, so Y is the current Closed flag."
        )
    doc.add_heading("Headline identification", level=1)
    doc.add_paragraph(
        f"CATE identified: {results.get('cate_identified')}. "
        f"Headline claim: {results.get('headline_claim')}."
    )
    add_bullets(doc, results.get("identification_reasons") or ["Identified."])
    return doc


def build_performance_report(
    results: dict[str, Any],
    figures: dict[str, Path] | None = None,
) -> Document:
    doc = new_doc(
        f"Uplift / CATE Performance Summary ({_title_suffix(results)})",
        "Milestone 6 — grouped multi-seed evaluation",
    )
    _notice(doc, _m7_notice(results))
    doc.add_heading("Sample", level=1)
    add_table(
        doc,
        ["metric", "value"],
        [
            {"metric": "Rows", "value": results.get("n_rows")},
            {"metric": "Members", "value": results.get("n_members")},
            {"metric": "Treated rows", "value": results.get("n_treated")},
            {"metric": "Control rows", "value": results.get("n_control")},
            {"metric": "Treated closure rate", "value": results.get("treated_closure_rate")},
            {"metric": "Control closure rate", "value": results.get("control_closure_rate")},
            {"metric": "ESS ratio", "value": results.get("ess_ratio")},
            {"metric": "CATE identified", "value": results.get("cate_identified")},
            {"metric": "Treatment grain", "value": results.get("treatment_grain")},
        ],
    )
    if results.get("per_seed"):
        doc.add_heading("Per-seed holdout", level=1)
        add_table(
            doc,
            ["seed", "ate", "auuc", "placebo_auuc"],
            results["per_seed"],
        )
    else:
        doc.add_paragraph(
            "CATE fitting was skipped because identification failed. "
            "Overlap diagnostics and the care-gap table are the deliverable."
        )
    if results.get("by_care_gap"):
        heading = (
            "Care-gap descriptives (Track R mock assignment)"
            if _is_mock_track_r(results)
            else "Care-gap descriptives (gap-level assignment, not CATE)"
        )
        doc.add_heading(heading, level=1)
        if _is_mock_track_r(results):
            doc.add_paragraph(
                "T is Bernoulli among gaps still open at reconstructed outreach. "
                "Y is 365-day post-issue InComm closure, not the raw Closed flag. "
                "Treated and control should both have outcome variation; T must not "
                "equal Y."
            )
        else:
            doc.add_paragraph(
                "Gap-level T follows InComm's live-validated rule: Closed ⇔ reward-eligible "
                "⇔ HEALTH_ACTION_REWARD case. Treated rows close at 100% and untreated at 0% "
                "because treatment is the outcome. Diabetes open rows are the untreated arm; "
                "the old untreated-empty diabetes table was member-broadcast contamination."
            )
        add_table(
            doc,
            [
                "care_gap_code",
                "n",
                "n_treated",
                "n_control",
                "treated_closure",
                "control_closure",
                "untreated_empty",
            ],
            results["by_care_gap"],
        )
    if results.get("positivity"):
        doc.add_heading("Positivity by stratum", level=1)
        doc.add_paragraph(
            "A stratum with no untreated (or treated) arm cannot support a tau estimate; none is reported for it, "
            "and the holdout size that would restore identification is listed."
        )
        add_table(
            doc,
            ["stratum", "n_treated", "n_control", "treated_rate", "control_rate", "status", "holdout_n_per_arm_for_1pp"],
            results["positivity"],
        )
    if results.get("followup_days"):
        fu = results["followup_days"]
        doc.add_paragraph(
            f"Follow-up: window {fu['window_days']} days; median available follow-up {fu['median']:.0f} days; "
            f"{100 * fu['share_full_window']:.1f}% of rows have the full window (the rest are censored at the data cut)."
        )
    for gap, sec in (results.get("secondary_cohorts") or {}).items():
        if sec.get("status") != "OK":
            continue
        doc.add_heading(f"Secondary cohort: {gap} (research-computed labels, not headline)", level=1)
        doc.add_paragraph(
            "InComm's rules engine never produced CKD rows (ConditionCode 'CKD' vs 'CHRONIC_KIDNEY_DISEASE'), so "
            "these labels are research-computed. Reported separately; never mixed into the headline."
        )
        add_table(
            doc,
            ["metric", "value"],
            [
                {"metric": "Rows", "value": sec.get("n_rows")},
                {"metric": "Treated / control", "value": f"{sec.get('n_treated')} / {sec.get('n_control')}"},
                {"metric": "Treated closure rate", "value": sec.get("treated_closure_rate")},
                {"metric": "Control closure rate", "value": sec.get("control_closure_rate")},
                {"metric": "ATE (SE)", "value": f"{sec.get('ate_mean')} ({sec.get('ate_se')})"},
            ],
        )
    grain = results.get("grain_contrast") or {}
    if grain.get("member_broadcast"):
        doc.add_heading("Why diabetes looked untreated-empty (member-broadcast T)", level=1)
        doc.add_paragraph(
            "M4 rolled sim rewards to MemberId and copied T onto every gap. "
            "A closed BH reward then marked still-open diabetes gaps as treated."
        )
        add_table(
            doc,
            [
                "care_gap_code",
                "n",
                "n_treated",
                "n_control",
                "treated_closure",
                "control_closure",
                "untreated_empty",
            ],
            grain["member_broadcast"],
        )
    if grain.get("notes"):
        doc.add_heading("Identification autopsy", level=1)
        add_bullets(doc, grain["notes"])
    _add_image(doc, figures.get("gap_contrast") if figures else None)
    _add_image(doc, figures.get("propensity") if figures else None)
    return doc


def build_elasticity_report(elasticity: dict[str, Any], results: dict[str, Any] | None = None) -> Document:
    results = results or {}
    doc = new_doc(
        f"Reward Elasticity Summary ({_title_suffix(results)})",
        "Milestone 6",
    )
    _notice(doc, _m7_notice(results) if results else f"Status: {elasticity.get('status')}")
    doc.add_paragraph(str(elasticity.get("reason", "")))
    if elasticity.get("pooled_slope_pp_per_10usd") is not None:
        add_table(
            doc,
            ["metric", "value"],
            [
                {"metric": "Pooled slope (pp closure per +$10 offered, within gap)", "value": elasticity["pooled_slope_pp_per_10usd"]},
                {"metric": "SE (pp per $10)", "value": elasticity["pooled_se_pp_per_10usd"]},
                {"metric": "Minimum detectable slope (80% power)", "value": elasticity["minimum_detectable_slope_pp_per_10usd"]},
                {"metric": "Distinguishable from zero", "value": elasticity["distinguishable_from_zero"]},
                {"metric": "Offers analysed", "value": elasticity["n_offered"]},
            ],
        )
    if elasticity.get("by_gap"):
        cols = ["care_gap_code", "n", "distinct_amounts", "mean_amount", "catalog_default"]
        if elasticity.get("pooled_slope_pp_per_10usd") is not None:
            cols += ["slope_pp_per_10usd", "se_pp_per_10usd"]
        add_table(doc, cols, elasticity["by_gap"])
    return doc


def build_targeting_report(
    ranked: pd.DataFrame,
    cate_identified: bool,
    results: dict[str, Any] | None = None,
    *,
    volume_tau: pd.DataFrame | None = None,
    volume_ablation: pd.DataFrame | None = None,
    figures: dict[str, Path] | None = None,
    ablation_importance: list[dict[str, Any]] | None = None,
) -> Document:
    results = results or {}
    doc = new_doc(
        f"Targeting Recommendation Output ({_title_suffix(results)})",
        "Milestone 6 — ranked by estimated uplift",
    )
    _notice(
        doc,
        "Decision-support list only. Not an automated denial of care. "
        f"{_m7_notice(results)} CATE identified: {cate_identified}.",
    )
    if ranked.empty:
        doc.add_paragraph("No targeting rows.")
        return doc
    if "quadrant" in ranked.columns:
        qcounts = ranked["quadrant"].value_counts().rename_axis("quadrant").reset_index(name="n")
        add_table(doc, ["quadrant", "n"], qcounts.to_dict(orient="records"))
    counts = ranked["recommended_band"].value_counts().rename_axis("band").reset_index(name="n")
    add_table(doc, ["band", "n"], counts.to_dict(orient="records"))
    _add_image(doc, figures.get("quadrant") if figures else None)
    if volume_tau is not None and not volume_tau.empty:
        doc.add_heading("Volume at τ thresholds", level=1)
        add_table(doc, ["threshold", "n_flagged", "flagged_share"], volume_tau.to_dict(orient="records"))
        _add_image(doc, figures.get("volume_tau") if figures else None)
    if volume_ablation is not None and not volume_ablation.empty:
        doc.add_heading("Volume at ablation P(close) thresholds (not uplift)", level=1)
        add_table(
            doc,
            ["threshold", "n_flagged", "flagged_share"],
            volume_ablation.to_dict(orient="records"),
        )
        _add_image(doc, figures.get("volume_ablation") if figures else None)
    if ablation_importance:
        doc.add_heading("Ablation model attributions (linear, no reward features)", level=1)
        add_table(
            doc,
            ["feature", "weight", "abs_weight"],
            ablation_importance[:12],
        )
    doc.add_heading("Top 25", level=1)
    headers = ["rank", "member_id", "care_gap_code", "uplift_score", "recommended_band"]
    if "quadrant" in ranked.columns:
        headers.append("quadrant")
    add_table(doc, headers, ranked.head(25).to_dict(orient="records"))
    return doc


def build_derived_layer_report(ddl: str, results: dict[str, Any] | None = None) -> Document:
    results = results or {}
    doc = new_doc(
        f"Derived-layer Integration Approach ({_title_suffix(results)})",
        "Proposed der_MemberRewardUpliftScores",
    )
    _notice(doc, _m7_notice(results) if results else "TreatmentSourceBasis is required on every row.")
    doc.add_paragraph(
        "The table mirrors der_MemberExperienceScores / der_MemberImpactScores. "
        "TreatmentSourceBasis is required on every row. Scores tagged "
        "Research-generated IncentiveRewards mock must not be read as InComm-issued."
    )
    doc.add_paragraph(ddl[:2000])
    return doc


def build_domain_resolution_report() -> Document:
    doc = new_doc(
        "Milestone 6 Domain-Analysis Resolutions",
        "Blockers closed per InComm direction to resolve via research, not wait",
    )
    _notice(doc, "These are research assumptions. They are not InComm production confirmations.")
    for key, item in DOMAIN_RESOLUTIONS.items():
        doc.add_heading(key, level=2)
        add_bullets(
            doc,
            [
                f"Blocker: {item['blocker']}",
                f"Resolution: {item['resolution']}",
                f"Safe claim: {item['claim']}",
            ],
        )
    return doc


def write_status_markdown(results: dict[str, Any], path: Path) -> str:
    path.parent.mkdir(parents=True, exist_ok=True)
    lines = [
        "# Milestone 6 Status",
        "",
        f"**Track:** {results.get('track', 'S')} ({_title_suffix(results)})",
        f"**TreatmentSourceBasis:** {results.get('treatment_source_basis')}",
        f"**Headline claim:** {results.get('headline_claim')}",
        f"**CATE identified:** {results.get('cate_identified')}",
        f"**Milestone 7 may consume:** {results.get('m7_may_consume', False)}",
        "",
        "## Identification",
        "",
    ]
    for reason in results.get("identification_reasons") or ["Identified."]:
        lines.append(f"- {reason}")
    lines += [
        "",
        "## Sample",
        "",
        f"- Rows: {results.get('n_rows')}",
        f"- Members: {results.get('n_members')}",
        f"- Treated closure: {results.get('treated_closure_rate')}",
        f"- Control closure: {results.get('control_closure_rate')}",
        "",
        "Do not treat Track S or research-generated Track R AUUC/ATE as partner-facing lift.",
        "",
    ]
    path.write_text("\n".join(lines), encoding="utf-8")
    return str(path)


def copy_shareables(docx_paths: list[Path], dest_dir: Path) -> list[str]:
    """Copy the six Word files into a Daniel-facing folder (same pattern as M5)."""
    dest_dir.mkdir(parents=True, exist_ok=True)
    copied: list[str] = []
    for src in docx_paths:
        target = dest_dir / src.name
        shutil.copy(src, target)
        copied.append(str(target))
    return copied


def write_all_reports(
    *,
    results: dict[str, Any],
    elasticity: dict[str, Any],
    ranked: pd.DataFrame,
    ddl: str,
    settings: Settings,
    figures: dict[str, Path] | None = None,
    volume_tau: pd.DataFrame | None = None,
    volume_ablation: pd.DataFrame | None = None,
    ablation_importance: list[dict[str, Any]] | None = None,
) -> list[str]:
    reports_dir = resolve_path(settings.paths.reports_dir)
    reports_dir.mkdir(parents=True, exist_ok=True)
    written: list[str] = []
    figures = figures or {}

    docs = {
        "Milestone 6 Reward Uplift Model v1.docx": build_method_card(results),
        "Milestone 6 Uplift Performance.docx": build_performance_report(results, figures),
        "Milestone 6 Elasticity Summary.docx": build_elasticity_report(elasticity, results),
        "Milestone 6 Targeting Recommendations.docx": build_targeting_report(
            ranked,
            bool(results.get("cate_identified")),
            results,
            volume_tau=volume_tau,
            volume_ablation=volume_ablation,
            figures=figures,
            ablation_importance=ablation_importance,
        ),
        "Milestone 6 Derived Layer Approach.docx": build_derived_layer_report(ddl, results),
        "Milestone 6 Domain Resolutions.docx": build_domain_resolution_report(),
    }
    docx_paths: list[Path] = []
    for name, doc in docs.items():
        dest = reports_dir / name
        doc.save(dest)
        written.append(str(dest))
        docx_paths.append(dest)

    share_dir = resolve_path(settings.milestone6.shareable_dir)
    written.extend(copy_shareables(docx_paths, share_dir))

    written.append(write_status_markdown(results, reports_dir / "Milestone6_Status.md"))
    payload = reports_dir / "milestone6_results.json"
    payload.write_text(json.dumps(results, indent=2, default=str), encoding="utf-8")
    written.append(str(payload))
    return written
