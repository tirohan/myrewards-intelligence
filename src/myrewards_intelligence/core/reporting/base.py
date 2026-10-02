"""Shared Word document generation helpers.

Generates professional deliverables for an external, non-technical-to-moderately-
technical reader: no internal file paths, no raw engineering details, no code-style
field names in table headers.
"""

from __future__ import annotations

from typing import Any

from docx import Document
from docx.shared import Pt

from ..logging import utc_now_iso

HEADER_LABELS: dict[str, str] = {
    "member_id": "Member ID",
    "care_gap_code": "Care Gap Code",
    "condition_code": "Condition",
    "status": "Status",
    "label_source": "Label Source",
    "feature_source": "Feature Source",
    "model_variant": "Model",
    "label_source_basis": "Label Source Basis",
    "metric": "Metric",
    "value": "Value",
    "threshold": "Threshold",
    "precision": "Precision",
    "recall": "Recall",
    "f1": "F1 Score",
    "auc": "AUC",
    "brier_score": "Brier Score",
    "subgroup": "Subgroup",
    "n_samples": "Sample Size",
    "n_positive": "Positive Cases",
    "base_rate": "Base Rate",
    "feature_name": "Feature",
    "importance": "Importance",
    "shap_value": "SHAP Value",
    "impact_score": "Impact Score",
    "score_tier": "Tier",
    "rank": "Rank",
    "description": "Description",
    "rationale": "Rationale",
    "n_treated": "Treated Rows",
    "n_control": "Control Rows",
    "treated_closure": "Treated Closure Rate",
    "control_closure": "Control Closure Rate",
    "untreated_empty": "Untreated Arm Empty",
    "uplift_score": "Uplift Score",
    "recommended_band": "Recommended Band",
    "quadrant": "Quadrant",
    "flagged_share": "Share Flagged",
    "n_flagged": "Rows Flagged",
    "feature": "Feature",
    "weight": "Weight",
    "abs_weight": "Absolute Weight",
}

_ACRONYMS = {"id", "auc", "shap", "lgbm", "utc", "ddl", "sql", "roi", "hedis"}


def _humanize(key: str) -> str:
    """Convert snake_case to Title Case, handling acronyms."""
    if key in HEADER_LABELS:
        return HEADER_LABELS[key]
    words = key.replace("_", " ").split()
    return " ".join(w.upper() if w.lower() in _ACRONYMS else w.capitalize() for w in words)


def _format_cell(value: Any) -> str:
    """Format a cell value for display."""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, float):
        if value != value:  # NaN: a value that is not available, never the string "nan"
            return "n/a"
        if abs(value) >= 1000:
            return f"{value:,.0f}"
        if abs(value) < 0.001:
            return f"{value:.4g}"
        return f"{value:.4f}"
    if value is None:
        return "—"
    return str(value)


def new_doc(title: str, subtitle: str = "") -> Document:
    """Create a new Word document with standard header."""
    doc = Document()
    doc.add_heading(title, level=0)
    if subtitle:
        p = doc.add_paragraph(subtitle)
        p.runs[0].italic = True
    date_para = doc.add_paragraph(f"Prepared by: KSU mHealth Research Lab   |   Date: {utc_now_iso()[:10]}")
    date_para.runs[0].font.size = Pt(9)
    return doc


def add_table(
    doc: Document,
    headers: list[str],
    rows: list[dict[str, Any]],
    max_rows: int = 400,
) -> None:
    """Add a formatted table to the document."""
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = _humanize(h)
        for p in hdr_cells[i].paragraphs:
            for r in p.runs:
                r.bold = True

    for row in rows[:max_rows]:
        cells = table.add_row().cells
        for i, h in enumerate(headers):
            cells[i].text = _format_cell(row.get(h, ""))

    if len(rows) > max_rows:
        note = doc.add_paragraph(
            f"(Showing {max_rows} of {len(rows)} rows. Complete table available on request.)"
        )
        note.runs[0].italic = True


def add_bullets(doc: Document, items: list[str]) -> None:
    """Add a bulleted list to the document."""
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_notice(doc: Document, text: str, bold: bool = True) -> None:
    """Add a prominent notice paragraph."""
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
